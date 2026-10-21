#!/usr/bin/env python3
"""
Genera data/calendario.js a partir del Excel CALENDARIO_SEMANAL.xlsx.

Uso:
    python scripts/actualizar_calendario.py RUTA/CALENDARIO_SEMANAL.xlsx
    python scripts/actualizar_calendario.py RUTA.xlsx --hoja "19 - 23 Octubre" --anio 2026
    python scripts/actualizar_calendario.py RUTA.xlsx --fecha 2026-10-21

Sin --hoja, busca la hoja de la semana que contiene --fecha (por defecto, hoy;
si es fin de semana, la semana siguiente). Las reglas de filtrado e importancia
se leen de config/calendario.json. Los datos publicados (campo "actual") se
conservan de data/calendario_actuales.json, que no se sobrescribe nunca.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent.parent
CONFIG = ROOT / "config" / "calendario.json"
ACTUALES = ROOT / "data" / "calendario_actuales.json"
OUT = ROOT / "data" / "calendario.js"

MESES = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
}
DIAS = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES"]
CODIGOS = {
    "ESPAÑA": "ES", "ITALIA": "IT", "ALEMANIA": "DE", "FRANCIA": "FR",
    "EUROZONA": "EZ", "EEUU": "US", "CHINA": "CN", "JAPÓN": "JP", "UK": "UK",
}
SHEET_RE = re.compile(r"^\s*(\d{1,2})\s*-\s*(\d{1,2})\s+([A-Za-zÁÉÍÓÚáéíóú]+)\s*$")


def semana_de_hoja(nombre: str, anio: int) -> dt.date | None:
    """'19 - 23 Octubre' -> lunes 2026-10-19. Devuelve None si no es una semana válida."""
    m = SHEET_RE.match(nombre)
    if not m:
        return None
    ini, fin, mes_txt = int(m.group(1)), int(m.group(2)), m.group(3)[:3].lower()
    mes = MESES.get(mes_txt)
    if not mes:
        return None
    try:
        fin_d = dt.date(anio, mes, fin)
    except ValueError:
        return None
    lunes = fin_d - dt.timedelta(days=4)
    return lunes if lunes.day == ini and lunes.weekday() == 0 else None


def elegir_hoja(wb, fecha: dt.date) -> tuple[str, dt.date]:
    objetivo = fecha - dt.timedelta(days=fecha.weekday())
    if fecha.weekday() >= 5:
        objetivo += dt.timedelta(days=7)
    for nombre in wb.sheetnames:
        lunes = semana_de_hoja(nombre, objetivo.year)
        if lunes == objetivo:
            return nombre, lunes
    sys.exit(f"No encuentro una hoja para la semana del {objetivo:%d/%m/%Y}. Usa --hoja.")


def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        return f"{v:g}".replace(".", ",")
    return str(v).strip()


def leer_hoja(ws) -> dict[str, list[dict]]:
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            v = getattr(c, "value", None)
            if v is not None:
                grid[(c.row, c.column)] = v

    # Cabeceras de día ('LUNES 19') y bloques ('Hora' + columnas a la derecha)
    dias = {(r, c): v for (r, c), v in grid.items()
            if isinstance(v, str) and v.split(" ")[0].upper() in DIAS}
    bloques = []
    for (r, c), v in grid.items():
        if v != "Hora":
            continue
        cols = {"hora": c}
        for dc in range(1, 13):
            h = grid.get((r, c + dc))
            key = {"País": "pais", "Dato": "dato", "Periodo": "periodo",
                   "Est.": "est", "Último": "ultimo"}.get(h)
            if key and key not in cols:
                cols[key] = c + dc
        # Día: la cabecera más cercana por encima en la misma columna
        cand = [(rr, v2) for (rr, cc), v2 in dias.items() if cc == c and rr < r]
        if not cand:
            continue
        dia = max(cand)[1].split(" ")[0].upper()
        bloques.append((dia, r, cols))

    # Si una cabecera 'Hora' aparece dos veces para el mismo día, nos quedamos con la última
    por_dia: dict[str, tuple[int, dict]] = {}
    for dia, r, cols in bloques:
        if dia not in por_dia or r > por_dia[dia][0]:
            por_dia[dia] = (r, cols)

    out: dict[str, list[dict]] = {}
    for dia, (r, cols) in por_dia.items():
        eventos = []
        rr = r + 1
        vacias = 0
        while vacias < 2 and rr < r + 80:
            pais = grid.get((rr, cols.get("pais")))
            dato = grid.get((rr, cols.get("dato")))
            if not pais and not dato:
                vacias += 1
                rr += 1
                continue
            vacias = 0
            eventos.append({
                "hora": fmt(grid.get((rr, cols["hora"]))),
                "pais": CODIGOS.get(str(pais).strip().upper(), str(pais).strip()),
                "dato": fmt(dato),
                "periodo": fmt(grid.get((rr, cols.get("periodo")))),
                "est": fmt(grid.get((rr, cols.get("est")))),
                "ultimo": fmt(grid.get((rr, cols.get("ultimo")))),
            })
            rr += 1
        out[dia] = eventos
    return out


def importancia(ev: dict, reglas: dict) -> str:
    texto = ev["dato"].lower()
    for nivel in ("alta", "media"):
        for patron in reglas.get("importancia", {}).get(nivel, []):
            p = patron.lower()
            if "|" in p:  # 'PAIS|texto'
                pais, txt = p.split("|", 1)
                if ev["pais"].lower() == pais and txt in texto:
                    return nivel
            elif p in texto:
                return nivel
    return "baja"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--hoja")
    ap.add_argument("--anio", type=int)
    ap.add_argument("--fecha", help="YYYY-MM-DD (por defecto hoy)")
    a = ap.parse_args()

    fecha = dt.date.fromisoformat(a.fecha) if a.fecha else dt.date.today()
    reglas = json.loads(CONFIG.read_text(encoding="utf-8")) if CONFIG.exists() else {}
    actuales = json.loads(ACTUALES.read_text(encoding="utf-8")) if ACTUALES.exists() else {}

    wb = openpyxl.load_workbook(a.excel, read_only=True, data_only=True)
    if a.hoja:
        lunes = semana_de_hoja(a.hoja, a.anio or fecha.year)
        if not lunes:
            sys.exit(f"No puedo deducir las fechas de la hoja '{a.hoja}'.")
        nombre = a.hoja
    else:
        nombre, lunes = elegir_hoja(wb, fecha)

    datos = leer_hoja(wb[nombre])
    paises = set(reglas.get("paises") or [])
    excluir = [x.lower() for x in reglas.get("excluir", [])]

    dias_out = []
    for i, dia in enumerate(DIAS):
        d = lunes + dt.timedelta(days=i)
        eventos = []
        for ev in datos.get(dia, []):
            if paises and ev["pais"] not in paises:
                continue
            if any(x in ev["dato"].lower() for x in excluir):
                continue
            ev["id"] = f'{d.isoformat()}|{ev["pais"]}|{ev["dato"]}'
            ev["imp"] = importancia(ev, reglas)
            a = actuales.get(ev["id"], "")
            if isinstance(a, dict):  # {"actual": ..., "est": ...}: est solo si el Excel no lo trae
                if not ev["est"] and a.get("est"):
                    ev["est"] = a["est"]
                    ev["est_fuente"] = "mercado"
                a = a.get("actual", "")
            ev["actual"] = a
            eventos.append(ev)
        eventos.sort(key=lambda e: e["hora"] or "99:99")
        dias_out.append({"fecha": d.isoformat(), "dia": dia.capitalize(), "eventos": eventos})

    payload = {
        "hoja": nombre,
        "semana": lunes.isoformat(),
        "generado": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "dias": dias_out,
    }
    OUT.write_text(
        "// Generado por scripts/actualizar_calendario.py — no editar a mano\n"
        "window.CALENDARIO = " + json.dumps(payload, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )
    n = sum(len(d["eventos"]) for d in dias_out)
    print(f"OK · hoja '{nombre}' · semana {lunes:%d/%m/%Y} · {n} eventos -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
